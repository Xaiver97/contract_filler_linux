# -*- coding: utf-8 -*-
"""
DOCX 导出 —— 把生成的合同图嵌入到 Word 文档中。

为什么需要这个模块：
    用户经常要把生成好的合同合并到一个文档里存档、转发或打印。
    这里把每张图嵌入到页面（A4 横向）作为纯图片段落。

要求（用户明确）：
    - 不要任何文字、标题、页眉页脚
    - 每页只放一张合同图，不留多余空页

设计要点：
    - 单户模式：一张图占一页
    - 批量模式：每户一页
    - 图片以 PNG 嵌入（保持清晰度）
    - 页面尺寸 A4 横向，图片自适应宽度
"""

import io
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image


def _image_to_png_bytes(img: Image.Image) -> tuple:
    """把 PIL 图转成 (png_bytes, width_px, height_px)。"""
    if img.mode != "RGB":
        img = img.convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    buf.seek(0)
    return buf, img.width, img.height


def _new_landscape_doc() -> Document:
    """新建一个 A4 横向文档，并清掉首段空白段落与页眉页脚内容。"""
    doc = Document()
    # 移除 python-docx 默认自带的空首段，避免顶部出现空白文字行
    for p in list(doc.paragraphs):
        if not p.text.strip():
            p._element.getparent().remove(p._element)

    # 清空页眉 / 页脚（不同首页、奇偶页都清），保证只留图片
    for section in doc.sections:
        section.different_first_page_header_footer = False
        for hf in (section.header, section.footer,
                   section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer):
            try:
                for para in list(hf.paragraphs):
                    for run in list(para.runs):
                        run._r.getparent().remove(run._r)
                    if para.text == "" and not para.runs:
                        para._p.getparent().remove(para._p)
            except Exception:
                pass

    section = doc.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = (
        section.page_height, section.page_width
    )
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    return doc


def _append_image_paragraph(doc: Document, img: Image.Image):
    """只把图片加成一个居中段落（无任何文字）。"""
    buf, w_px, h_px = _image_to_png_bytes(img)
    usable_w_cm = 27.7  # A4 横向可用宽度(cm)，约 29.7 - 1.0*2
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(buf, width=Cm(usable_w_cm))
    return doc


def export_one(img: Image.Image, path) -> Path:
    """导出单户合同为 DOCX：一页纯图片，无文字/页眉/页脚。"""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    doc = _new_landscape_doc()
    _append_image_paragraph(doc, img)
    doc.save(str(p))
    return p


def export_many(items: list, path, progress_cb=None) -> Path:
    """批量导出多户合同为单个 DOCX：每页一张纯图片。

    参数
    ----
    items : list of (caption, image)    兼容旧接口，caption 被忽略（不出文字）
    path : str | Path
    progress_cb : callable(i, total)   可选进度回调
    """
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)

    doc = _new_landscape_doc()
    n = len(items)
    for i, (_caption, img) in enumerate(items):
        if i > 0:
            doc.add_page_break()   # 仅在两张之间分页，避免末尾空页
        _append_image_paragraph(doc, img)
        if progress_cb:
            progress_cb(i + 1, n)
    doc.save(str(p))
    return p