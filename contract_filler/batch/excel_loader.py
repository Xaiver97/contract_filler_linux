# -*- coding: utf-8 -*-
"""
Excel 读写 —— 批量导入与模板导出。

设计要点：
    1. 表头宽容匹配：中文标签、英文 key、常见别名都认
    2. 类型清洗：Excel 的 12 可能读成 12.0，日期可能读成 datetime
    3. 校验：可选值字段（用电容量/电压等级等）越界时给出警告而非报错
"""

import datetime as _dt
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# 表头别名：key -> 可识别的中文/英文写法
ALIASES = {
    "customer_name": ["用电人", "用户名称", "客户名称", "户名", "用户"],
    "address":       ["用电地址", "地址", "用电地点", "装表地址"],
    "usage_nature":  ["用电性质", "性质", "用电类别"],
    "capacity":      ["用电容量", "容量", "合同容量", "受电容量"],
    "voltage":       ["电压等级", "电压"],
    "customer_no":   ["客户编号", "户号", "编号", "用户编号"],
    "limit_plan":    ["限额方案", "限额"],
    "phone":         ["手机号码", "手机号", "电话", "联系方式", "联系电话"],
    "signature":     ["用电人签名", "签名", "签字", "签章"],
    "id_card":       ["身份证号码", "身份证号", "身份证"],
    "sign_date":     ["签订日期", "日期", "签约日期", "签订时间"],
    "notice_text":   ["已知晓声明", "告知内容", "声明", "备注声明"],
}

# 必须按文本处理的列。
# 关键：18 位身份证号若被 Excel 存成数字，float64 只有约 15 位有效数字，
# 后三位会被抹成 0 —— 这是批量填表最隐蔽的坑，必须在模板阶段就防住。
TEXT_KEYS = {"id_card", "phone", "customer_no", "sign_date", "capacity",
             "limit_plan"}

# 超过这个量级的整数，一旦以 float 形式进出 Excel 就可能丢精度
BIG_NUMBER = 1e15


def _norm(s):
    return str(s).strip().replace(" ", "").replace("\u3000", "").lower()


def _match_key(header, fields):
    """把表头文字匹配到字段 key。匹配不上返回 None。"""
    if not header:
        return None
    h = _norm(header)

    for f in fields:
        key = f.get("key", "")
        if not key:
            continue
        if h == _norm(key):
            return key
        label = f.get("label", "")
        if label and h == _norm(label):
            return key
        for alias in ALIASES.get(key, []):
            if h == _norm(alias):
                return key
    # 再退一步：包含匹配
    for f in fields:
        key = f.get("key", "")
        label = f.get("label", "")
        for cand in (label, key):
            if cand and _norm(cand) and _norm(cand) in h:
                return key
    return None


def _clean_value(v, date_format="%Y.%m.%d"):
    """把单元格原始值清洗成可直接书写的字符串。"""
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime(date_format)
    if isinstance(v, float):
        # 12.0 -> "12"，避免合同上写出 "12.0"
        if abs(v - round(v)) < 1e-9:
            return str(int(round(v)))
        return str(v)
    if isinstance(v, int):
        return str(v)
    return str(v).strip()


def load_excel(path, fields, date_format="%Y年%m月%d日",
               sheet=None, has_header=True):
    """读取 Excel，返回 (记录列表, 警告列表)。

    记录是 dict，键为字段 key，另含特殊键：
        _row        : int（Excel 中的行号，便于报错定位）
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"找不到 Excel 文件: {path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return [], ["表格是空的"]

    warnings = []
    if has_header:
        header = rows[0]
        data_rows = rows[1:]
    else:
        header = None
        data_rows = rows

    colmap = {}
    if has_header:
        used = {}
        for idx, h in enumerate(header):
            if h is None:
                continue
            if str(h).strip().startswith("#"):
                continue
            key = _match_key(h, fields)
            if key:
                if key in used:
                    warnings.append(
                        f"表头重复：「{h}」与「{used[key]}」都映射到 {key}，"
                        f"以首次出现的列为准")
                    continue
                used[key] = h
                colmap[idx] = key
        # 注意：colmap 是 {列索引: key}，这里要查的是 value 而不是 key
        mapped = set(colmap.values())
        missing = [f.get("label", f.get("key"))
                   for f in fields
                   if f.get("required") and f.get("key") not in mapped]
        if missing:
            warnings.append("缺少必填列：" + "、".join(missing))
        if not colmap:
            warnings.append("未能识别任何表头，请检查列名是否与模板一致")
    else:
        for idx, f in enumerate(fields):
            colmap[idx] = f.get("key")

    # 可选值约束
    option_map = {}
    for f in fields:
        opts = f.get("options")
        if f.get("type") in ("checkbox_group", "radio"):
            opts = [o.get("value") for o in f.get("options", [])]
        if opts:
            option_map[f.get("key")] = [str(o) for o in opts]

    records = []
    for r_idx, row in enumerate(data_rows):
        if row is None:
            continue
        if all(c is None or str(c).strip() == "" for c in row):
            continue

        rec = {"_row": (r_idx + 2) if has_header else (r_idx + 1)}
        for c_idx, key in colmap.items():
            if c_idx >= len(row):
                continue
            raw = row[c_idx]
            # 长号码以数字形式存储 → 精度已不可恢复，必须提醒用户
            if (isinstance(raw, (int, float))
                    and abs(float(raw)) >= BIG_NUMBER):
                warnings.append(
                    f"第 {rec['_row']} 行：{key} 在 Excel 里是数字，"
                    f"长号码会丢失精度（{raw}）。请把该列设为「文本」格式，"
                    f"或在输入前加英文单引号。")
            rec[key] = _clean_value(raw, date_format)

        # 可选值校验
        for key, allowed in option_map.items():
            val = rec.get(key)
            if val not in (None, "") and str(val) not in allowed:
                warnings.append(
                    f"第 {rec['_row']} 行：{key} 的值「{val}」不在可选范围 "
                    f"({'、'.join(allowed)})，仍会照原样写入")

        records.append(rec)

    return records, warnings


def export_template(path, fields, sample: dict = None, with_sample=True):
    """导出一份带表头、示例行与下拉验证的 Excel 模板。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "合同数据"

    header_fill = PatternFill("solid", fgColor="DCE6F1")
    header_font = Font(bold=True)

    cols = list(fields)

    # 表头
    for c, f in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=f.get("label", f.get("key")))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = 18

    # 下拉验证：用区域字符串而不是逐单元格 add。
    # 逐单元格 add 会预建出几百个空行，导致用户追加数据时行号错位。
    for c, f in enumerate(cols, start=1):
        opts = None
        if f.get("type") in ("select",):
            opts = [str(o) for o in f.get("options", [])]
        elif f.get("type") in ("checkbox_group", "radio"):
            opts = [str(o.get("value")) for o in f.get("options", [])]
        if opts:
            dv = DataValidation(type="list",
                                formula1='"' + ",".join(opts) + '"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            letter = get_column_letter(c)
            dv.add(f"{letter}2:{letter}2000")

    # 示例行
    if with_sample:
        demo = sample or {}
        for c, f in enumerate(cols, start=1):
            key = f.get("key")
            val = demo.get(key, f.get("default", ""))
            ws.cell(row=2, column=c, value=val)

    # 长号码列设为文本格式，防止身份证等被 Excel 吃掉精度。
    # 整列样式打底，再确保示例行单元格本身也是文本。
    for c, f in enumerate(cols, start=1):
        if f.get("key") in TEXT_KEYS:
            letter = get_column_letter(c)
            ws.column_dimensions[letter].number_format = "@"
            ws.cell(row=2, column=c).number_format = "@"

    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def demo_row(fields):
    """生成一行示例数据，供导出模板使用。"""
    demo = {
        "customer_name": "张伟",
        "address": "江西省南昌市红谷滩区丰和大道 1368 号",
        "usage_nature": "一般工商业",
        "capacity": "39",
        "voltage": "220",
        "customer_no": "2520471",
        "limit_plan": "1",
        "phone": "13800000000",
        "notice_text": "已知晓自动停电规则，一欠费就停电",
        "signature": "张伟",
        "id_card": "360101199001011234",
        "sign_date": _dt.date.today().strftime("%Y.%m.%d"),
    }
    return {k: v for k, v in demo.items() if any(
        f.get("key") == k for f in fields)}
